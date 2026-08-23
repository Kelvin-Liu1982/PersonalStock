#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
应收账款增速差：计算与回填
========================
依据：应收账款增速差_指标说明.md 的计算公式

    广义应收账款   = 应收账款 + 应收票据 + 合同资产
    广义应收增速%  = (本期广义应收 − 上期广义应收) / 上期广义应收 × 100
    营收增速%      = (本期营收 − 上期营收) / 上期营收 × 100
    应收账款增速差  = 广义应收增速% − 营收增速%          （单位：百分点 pp）

本脚本：
  1) 定义可复用的 calc_receivable_growth_gap() 函数（不依赖 Excel，纯计算）；
  2) 从「中国移动_财务健康与盈利质量.xlsx」年报数据 sheet 读取
     应收账款 / 应收票据 / 合同资产 / 营业收入；
  3) 计算后在 sheet 中新增两列：
        · 广义应收账款(亿元,年末净额)
        · 应收账款增速差(百分点pp)
     并在首行（无同比基期）标注 “—”。
  4) 幂等：若目标列已存在，仅刷新数值，不重复插入。
"""
import openpyxl

XLSX = "中国移动_财务健康与盈利质量.xlsx"
SHEET = "年报数据"


def calc_receivable_growth_gap(years, ar, bn, ca, rev):
    """计算应收账款增速差。

    参数（均按年份升序、等长序列，金额单位需一致）：
        years : 年份列表，如 [2020, 2021, ...]
        ar    : 应收账款（净额）
        bn    : 应收票据（净额）
        ca    : 合同资产
        rev   : 营业收入
    返回：与 years 等长的 list[dict]，每项含：
        year, broad_ar(广义应收), broad_ar_yoy, rev_yoy, gap_pp
        首年无同比基期，broad_ar_yoy / rev_yoy / gap_pp 为 None。
    """
    n = len(years)
    broad = [ar[i] + bn[i] + ca[i] for i in range(n)]
    out = []
    for i in range(n):
        row = {"year": years[i], "broad_ar": round(broad[i], 2)}
        if i == 0:
            row.update({"broad_ar_yoy": None, "rev_yoy": None, "gap_pp": None})
        else:
            broad_yoy = (broad[i] - broad[i - 1]) / broad[i - 1] * 100
            rev_yoy = (rev[i] - rev[i - 1]) / rev[i - 1] * 100
            row.update({
                "broad_ar_yoy": round(broad_yoy, 2),
                "rev_yoy": round(rev_yoy, 2),
                "gap_pp": round(broad_yoy - rev_yoy, 2),
            })
        out.append(row)
    return out


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]

    # 列位置：A 年份(1), B 营收(2), O 应收(15), P 票据(16), Q 合同资产(17)
    years, ar, bn, ca, rev = [], [], [], [], []
    for r in range(5, ws.max_row + 1):
        y = ws.cell(r, 1).value
        if y is None:
            continue
        years.append(y)
        ar.append(float(ws.cell(r, 15).value))
        bn.append(float(ws.cell(r, 16).value))
        ca.append(float(ws.cell(r, 17).value))
        rev.append(float(ws.cell(r, 2).value))

    results = calc_receivable_growth_gap(years, ar, bn, ca, rev)

    # 幂等：若已存在目标列，跳过插入
    header_row = 4
    existing = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    if "应收账款增速差(百分点pp)" not in existing:
        # 在原 R 列（应收/营收匹配）之前插入 2 列
        ws.insert_cols(18, 2)
        ws.cell(header_row, 18, "广义应收账款(亿元,年末净额)")
        ws.cell(header_row, 19, "应收账款增速差(百分点pp)")
        from openpyxl.styles import Font
        src = ws.cell(header_row, 17).font
        fnt = Font(name=src.name, size=src.size, bold=src.bold,
                   italic=src.italic, color=src.color)
        ws.cell(header_row, 18).font = fnt
        ws.cell(header_row, 19).font = fnt
        ws.column_dimensions["R"].width = 22
        ws.column_dimensions["S"].width = 22

    for i, res in enumerate(results):
        r = 5 + i
        ws.cell(r, 18, res["broad_ar"]).number_format = "0.00"
        gap = res["gap_pp"]
        ws.cell(r, 19, "—" if gap is None else f"{gap:+.2f}pp")

    wb.save(XLSX)

    # 控制台校验输出
    print(f"{'年份':>6} {'广义应收(亿)':>12} {'广义应收增速%':>13} {'营收增速%':>10} {'增速差(pp)':>12}")
    print("-" * 56)
    for res in results:
        gy = "   —   " if res["broad_ar_yoy"] is None else f"{res['broad_ar_yoy']:+.2f}"
        ry = "   —   " if res["rev_yoy"] is None else f"{res['rev_yoy']:+.2f}"
        gp = "   —   " if res["gap_pp"] is None else f"{res['gap_pp']:+.2f}"
        print(f"{res['year']:>6} {res['broad_ar']:>12.2f} {gy:>13} {ry:>10} {gp:>12}")


if __name__ == "__main__":
    main()
