# -*- coding: utf-8 -*-
"""生成卡片 4（客户生态 CHBN）7 个指标的年度 SVG 折线图。

数据源：数据源/中国移动_客户生态CHBN.xlsx · 年报数据 sheet
输出：svg 图/s4/s4_*.svg
仅读取 XLSX，绝不回写。
"""
import os, re, sys
ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, "/Users/sto1/.workbuddy/skills/finance-xlsx-svg-charts/scripts")
sys.path.insert(0, ROOT)
import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from common.svg_charts import make_line_svg
import openpyxl

XLSX = os.path.join(ROOT, "数据源", "中国移动_客户生态CHBN.xlsx")
OUT = os.path.join(ROOT, "svg 图", "s4")
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["年报数据"]

# R4 表头，找各指标列号
header = {}
for c in range(1, ws.max_column + 1):
    v = ws.cell(4, c).value
    if v:
        header[str(v).strip()] = c

# 年份行：R5=2020 ... R10=2025；取 2021-2025 (R6..R10)
year_rows = {2020: 5, 2021: 6, 2022: 7, 2023: 8, 2024: 9, 2025: 10}
want = [2021, 2022, 2023, 2024, 2025]

def parse(v, is_ratio=False):
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "—", "-", "约", "未单独披露"):
        return None
    num = re.search(r"-?\d+\.?\d*", s)
    if not num:
        return None
    x = float(num.group())
    if is_ratio and abs(x) < 1:      # 小数比例 → 百分比
        x *= 100
    return x

# 7 指标：(key, 表头名, 单位, 颜色, 是否比例)
specs = [
    ("c",     "个人市场C（亿）", "亿元",   "#c8102e", False),
    ("h",     "家庭市场H（亿）", "亿元",   "#2e9e6b", False),
    ("b",     "政企市场B（亿）", "亿元",   "#15616d", False),
    ("n",     "新兴市场N（亿）", "亿元",   "#d98a14", False),
    ("hbn",   "HBN占通服比",     "%",      "#c8102e", True),
    ("corp",  "政企客户数（万家）", "万家", "#15616d", False),
    ("cloud", "移动云收入（亿）", "亿元",  "#2e9e6b", False),
]

YEARS = [str(y) for y in want]
for key, hname, unit, color, is_ratio in specs:
    col = header[hname]
    vals = [parse(ws.cell(year_rows[y], col).value, is_ratio) for y in want]
    title = hname.split("（")[0]
    svg = make_line_svg(title, unit, YEARS, vals, color, W=480, H=260)
    fn = os.path.join(OUT, f"s4_{key}_annual.svg")
    with open(fn, "w", encoding="utf-8") as f:
        f.write(svg)
    print(f"OK  s4_{key}_annual.svg  {title}  values={vals}")

print("\n生成完成，输出目录：", OUT)
