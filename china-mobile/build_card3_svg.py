# -*- coding: utf-8 -*-
"""生成卡片 3（资本开支与基础设施）6 个指标的年度 SVG 折线图。

数据源：数据源/中国移动_资本开支与基础设施.xlsx · 年报数据 sheet
输出：svg 图/s3/s3_*.svg
仅读取 XLSX，绝不回写。
"""
import os, sys
ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, "/Users/sto1/.workbuddy/skills/finance-xlsx-svg-charts/scripts")
sys.path.insert(0, ROOT)
from gen_line_svg import make_line_svg

import openpyxl

XLSX = os.path.join(ROOT, "数据源", "中国移动_资本开支与基础设施.xlsx")
OUT = os.path.join(ROOT, "svg 图", "s3")
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["年报数据"]

# R4 表头：指标 | 2020 | 2021 | 2022 | 2023 | 2024 | 2025 | 2026F/备注
data = {}  # name -> [2020..2025]
for r in range(5, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name:
        continue
    vals = []
    for c in range(2, 8):  # 2020..2025
        v = ws.cell(r, c).value
        if v in (None, "", "—", "-", "约"):
            vals.append(None)
        else:
            try:
                vals.append(float(v))
            except (TypeError, ValueError):
                vals.append(None)
    data[name.strip()] = vals

def slice2021(v2020_2025):
    # 取 2021..2025（去掉 2020），与趋势区标题 2021–2026F 对齐
    return v2020_2025[1:6]

XL = "—"  # 占位
YEARS = ["2021", "2022", "2023", "2024", "2025", "2026"]

# 6 个卡片指标： (文件key, xlsx指标名, 单位, 颜色, 2026值, 是否H1实际)
specs = [
    ("capex",    "资本开支总额（亿元）",        "亿元",    "#15616d", 1366,  False),
    ("net",      "通信网络投资（亿元）",        "亿元",    "#7a8694", 792,   False),
    ("compute",  "算力网络投资（亿元）",        "亿元",    "#2e9e6b", 378,   False),
    ("aiinv",    "智能网络投资（亿元）",        "亿元",    "#d98a14", 89,    False),
    ("ai",       "智算总规模（EFLOPS, FP16）",  "EFLOPS",  "#c8102e", 112.7, True),
    ("5g",       "5G基站累计（万站）",          "万站",    "#15616d", 294,   True),
]

for key, xname, unit, color, v2026, is_h1 in specs:
    v = slice2021(data[xname])[:]          # [2021..2025]
    v.append(v2026)                        # 追加 2026
    # 指标名去单位后缀做标题
    title = xname.split("（")[0]
    svg = make_line_svg(title, unit, YEARS, v, color, W=480, H=260)
    fn = os.path.join(OUT, f"s3_{key}_annual.svg")
    with open(fn, "w", encoding="utf-8") as f:
        f.write(svg)
    print(f"OK  s3_{key}_annual.svg  {title}  values={v}  (2026={'H1' if is_h1 else 'F'})")

print("\n生成完成，输出目录：", OUT)
