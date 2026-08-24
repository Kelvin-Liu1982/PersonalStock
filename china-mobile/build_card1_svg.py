# -*- coding: utf-8 -*-
"""读取 数据源/《中国移动_财务健康与盈利质量.xlsx》，为看板卡片1（财务健康与盈利质量）
的指标生成年度 / 季度 SVG 折线图（纯静态 SVG，无 JS 依赖）。

营业收入、归母净利润、EBITDA率 有季度序列（2020Q1–2026Q2 单季），生成季度 svg；
ROE、自由现金流、应收账款增速差 仅年报披露，仅生成年度图。

输入： 数据源/中国移动_财务健康与盈利质量.xlsx（年报数据 / 季报数据 两张 sheet）
输出： svg 图/s1/ 下
      年度 6 张：s1_rev/np/eb/roe/fcf/gap_annual.svg
      季度 3 张：s1_rev/np/eb_q.svg（EBITDA率=EBITDA/营收）
"""
import openpyxl
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from gen_line_svg import make_line_svg, make_placeholder

FN = os.path.join(HERE, "数据源", "中国移动_财务健康与盈利质量.xlsx")
OUT_DIR = os.path.join(HERE, "svg 图", "s1")
os.makedirs(OUT_DIR, exist_ok=True)

# ---------- 读取 XLSX ----------
wb = openpyxl.load_workbook(FN, data_only=True)

# 年报数据
wa = wb["年报数据"]
hdr_a = [wa.cell(4, c).value for c in range(1, 21)]
hidx_a = {h: i for i, h in enumerate(hdr_a) if h}
def ann(header):
    ci = hidx_a[header] + 1
    return [wa.cell(r, ci).value for r in range(5, 11)]

years = [2020, 2021, 2022, 2023, 2024, 2025]
rev_a = ann("营业收入(亿元)")
np_a = ann("归母净利润(亿元)")
eb_a = ann("EBITDA率(%)")
roe_a = ann("ROE(%)")
fcf_a = ann("自由现金流(亿元)")
gap_raw = ann("应收账款增速差(百分点pp)")
gap_a = []
for v in gap_raw:
    if isinstance(v, str) and "—" in v:
        gap_a.append(None)
    elif isinstance(v, (int, float)):
        gap_a.append(float(v))
    elif isinstance(v, str):
        gap_a.append(float(v.replace("pp", "").replace("+", "")))
    else:
        gap_a.append(None)

ANN = [
    ("rev", "营业收入", "亿元", rev_a, "#c8102e"),
    ("np", "归母净利润", "亿元", np_a, "#15616d"),
    ("eb", "EBITDA率", "%", eb_a, "#2e9e6b"),
    ("roe", "ROE（摊薄）", "%", roe_a, "#0e4a55"),
    ("fcf", "自由现金流", "亿元", fcf_a, "#7a8694"),
    ("gap", "应收账款增速差", "pp", gap_a, "#d98a14"),
]

# ---------- 季报数据（仅 rev / np / eb 有季度序列）----------
wq = wb["季报数据"]
hdr_q = [wq.cell(4, c).value for c in range(1, 9)]
qidx_q = {h: i for i, h in enumerate(hdr_q) if h}
quarters = [wq.cell(r, 1).value for r in range(5, 31)]
def qv(header):
    ci = qidx_q[header] + 1
    return [wq.cell(r, ci).value for r in range(5, 31)]

q_rev = qv("营业收入(亿元)")
q_np = qv("归母净利润(亿元)")
q_ebitda = qv("EBITDA(亿元)")
q_eb = []
for e, r in zip(q_ebitda, q_rev):
    if isinstance(e, (int, float)) and isinstance(r, (int, float)) and r:
        q_eb.append(e / r * 100)
    else:
        q_eb.append(None)

Q_HAS = [
    ("rev", "营业收入", "亿元", q_rev, "#c8102e"),
    ("np", "归母净利润", "亿元", q_np, "#15616d"),
    ("eb", "EBITDA率", "%", q_eb, "#2e9e6b"),
]

# ---------- 生成并写文件 ----------
written = []
for key, name, unit, vals, color in ANN:
    svg = make_line_svg(name, unit, years, vals, color, W=480, H=260, quarterly=False)
    fn = os.path.join(OUT_DIR, f"s1_{key}_annual.svg")
    with open(fn, "w", encoding="utf-8") as f:
        f.write(svg)
    written.append(fn)

for key, name, unit, vals, color in Q_HAS:
    svg = make_line_svg(name, unit, quarters, vals, color, W=720, H=280, quarterly=True)
    fn = os.path.join(OUT_DIR, f"s1_{key}_q.svg")
    with open(fn, "w", encoding="utf-8") as f:
        f.write(svg)
    written.append(fn)

print("生成 %d 个 SVG 文件（输出至 %s）：" % (len(written), os.path.relpath(OUT_DIR, HERE)))
for w in written:
    print("  ", os.path.basename(w))

# 回显关键数值，便于核对
print("\n年度核对：")
for key, name, unit, vals, color in ANN:
    print(f"  {name}: {vals}")
print("\n季度（首末点）核对：")
for key, name, unit, vals, color in Q_HAS:
    print(f"  {name}: n={len(vals)} 首={vals[0]} 末={vals[-1]}")
