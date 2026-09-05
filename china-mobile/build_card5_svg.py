# -*- coding: utf-8 -*-
"""生成卡片 5（宏观政策与股东回报）前 4 个数据指标的年度 SVG 折线图。

数据源：数据源/中国移动_宏观政策与股东回报.xlsx · 年报数据 sheet
输出：svg 图/s5/s5_*.svg
仅读取 XLSX，绝不回写。
4 指标 = 卡片 5 的 4 个 sparkline 系列：派息率 / 每股股息 / 股息率(H股) / 股息利差
（跳过 R8 中国10年期国债收益率，因卡片 5 未将其列为数据指标）。
x 轴统一 2020–2025，与卡片内 sk5_* sparkline 取值对齐。
"""
import os, re, sys
ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, "/Users/sto1/.workbuddy/skills/finance-xlsx-svg-charts/scripts")
sys.path.insert(0, ROOT)
import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from common.svg_charts import make_line_svg
import openpyxl

XLSX = os.path.join(ROOT, "数据源", "中国移动_宏观政策与股东回报.xlsx")
OUT = os.path.join(ROOT, "svg 图", "s5")
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["年报数据"]

# R4 表头：指标 | 2020 | 2021 | ... | 2025 | 备注
# 年份列号映射
year_col = {}
for c in range(1, ws.max_column + 1):
    v = ws.cell(4, c).value
    if v and str(v).strip() in ("2020", "2021", "2022", "2023", "2024", "2025"):
        year_col[str(v).strip()] = c

# 指标名（第1列）→ 行号
name_row = {}
for r in range(5, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v:
        name_row[str(v).strip()] = r

YEARS = ["2020", "2021", "2022", "2023", "2024", "2025"]


def parse(v):
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "—", "-", "约", "未单独披露"):
        return None
    num = re.search(r"-?\d+\.?\d*", s)
    if not num:
        return None
    return float(num.group())


# 4 指标：(key, 指标名, 单位, 颜色)
# key 须与卡片 5 sparkline id 后缀一致：sk5_<key> → s5_<key>_annual
specs = [
    ("payout", "派息率（%）", "%", "#2e9e6b"),
    ("dps",    "每股股息（港元）", "港元", "#c8102e"),
    ("yield",  "H股股息率", "%", "#d98a14"),
    ("spread", "股息利差（百分点）", "pp", "#15616d"),
]

for key, hname, unit, color in specs:
    r = name_row[hname]
    vals = [parse(ws.cell(r, year_col[y]).value) for y in YEARS]
    title = hname.split("（")[0]
    svg = make_line_svg(title, unit, YEARS, vals, color, W=480, H=260)
    fn = os.path.join(OUT, f"s5_{key}_annual.svg")
    with open(fn, "w", encoding="utf-8") as f:
        f.write(svg)
    print(f"OK  s5_{key}_annual.svg  {title}  values={vals}")

print("\n生成完成，输出目录：", OUT)
