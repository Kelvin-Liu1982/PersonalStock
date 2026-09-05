# -*- coding: utf-8 -*-
"""
只读 数据源/中煤能源_产销数据.xlsx，生成「自产商品煤产量」与「商品煤总销量」
的【季度】+【年度】趋势图、「自产/贸易煤结构·自产占比」趋势图，以及卡片内 sparkline；
并读取「在产矿井核定产能(年度)」sheet 生成在产产能年度趋势图、
读取「剩余可采储量(年度)」sheet 生成剩余可采储量年度趋势图（亿吨，绿）。
另只读 数据源/中煤能源_价格和成本.xlsx（原秦皇岛5500现货价_周度.xlsx 改名），
读取「自产煤平均售价(年度)」「自产煤吨完全成本(年度)」「自产煤吨毛利(年度)」三 sheet
生成对应的年度趋势图（trend_/spark_·元/吨；售价红 / 成本绿 / 毛利红）。

重要：本脚本只读取 xlsx、只写出 SVG，**绝不写回 xlsx**（xlsx 由用户手工维护，
含 2020-2026 季度明细与 2020-2025 全年合计；2026 全年进行中、无全年合计）。
旧的 build_coal_volume.py 因含硬编码年度值且会 wb.save() 覆盖 xlsx，已删除，
其年度趋势图功能由本脚本的 build_trend_y 统一接管。

用法（在 china-coal/ 目录下）：
    python3 build_volume_svg.py
依赖：openpyxl（仅读取）
"""
import os
import openpyxl

SVG_DIR = "svg"   # 图表统一输出到 svg/ 子目录
os.makedirs(SVG_DIR, exist_ok=True)

FID = "数据源/中煤能源_产销数据.xlsx"
PR_FID = "数据源/中煤能源_价格和成本.xlsx"  # 价格与成本（原秦皇岛5500现货价_周度.xlsx 改名）

# ---------- 读取：从两个 sheet 提取 2020Q1-2026Q2 季度序列 ----------
def extract_quarterly(ws):
    hdr = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "年份":
            hdr = r
            break
    if hdr is None:
        raise ValueError("未找到表头行（年份）")
    series = []
    for r in range(hdr + 1, ws.max_row + 1):
        y = ws.cell(row=r, column=1).value
        if not isinstance(y, (int, float)):
            continue  # 跳过汇总/说明行
        y = int(y)
        for q in range(1, 5):
            v = ws.cell(row=r, column=1 + q).value
            if isinstance(v, (int, float)):
                series.append((y, q, float(v)))
    return series


# ---------- 读取：从 全年合计 列提取年度序列（2020-2025）----------
def extract_annual(ws):
    hdr = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "年份":
            hdr = r
            break
    if hdr is None:
        raise ValueError("未找到表头行（年份）")
    col = None
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=hdr, column=c).value
        if h and "全年合计" in str(h):
            col = c
            break
    if col is None:
        raise ValueError("未找到『全年合计』列")
    series = []
    for r in range(hdr + 1, ws.max_row + 1):
        y = ws.cell(row=r, column=1).value
        if not isinstance(y, (int, float)):
            continue
        v = ws.cell(row=r, column=col).value
        if isinstance(v, (int, float)):   # 跳过无全年合计的行（如 2026 进行中）
            series.append((int(y), float(v)))
    return series


# ---------- 读取：按表头子串定位某列，提取年度序列（通用）----------
def extract_annual_named(ws, header_substr):
    hdr = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "年份":
            hdr = r
            break
    if hdr is None:
        raise ValueError("未找到表头行（年份）")
    col = None
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=hdr, column=c).value
        if h and header_substr in str(h):
            col = c
            break
    if col is None:
        raise ValueError(f"未找到含『{header_substr}』的列")
    series = []
    for r in range(hdr + 1, ws.max_row + 1):
        y = ws.cell(row=r, column=1).value
        if not isinstance(y, (int, float)):
            continue
        v = ws.cell(row=r, column=col).value
        if isinstance(v, (int, float)):
            series.append((int(y), float(v)))
    return series


# ---------- 季度趋势图（对齐 trend_revenue_q.svg 风格）----------
def build_trend_q(series, fname, title, color="#c53030"):
    W, H = 1000, 420
    ml, mr, mt = 58, 24, 42
    plot_bottom = 376
    ph = plot_bottom - mt          # 334
    pw = W - ml - mr               # 918
    n = len(series)
    xs = [ml + pw * i / (n - 1) for i in range(n)]
    vals = [s[2] for s in series]
    vmin, vmax = min(vals), max(vals)
    pad = (vmax - vmin) * 0.18 or 10
    vmin, vmax = vmin - pad, vmax + pad
    def yp(v):
        return mt + ph - (v - vmin) / (vmax - vmin) * ph
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" '
           f'font-family="-apple-system,Segoe UI,PingFang SC,sans-serif">']
    svg.append(f'<rect width="{W}" height="{H}" fill="#ffffff"/>')
    svg.append(f'<text x="{ml}" y="26" font-size="16" font-weight="700" fill="#1e3a5f">{title}</text>')
    for k in range(5):
        yv = vmin + (vmax - vmin) * k / 4
        yy = yp(yv)
        svg.append(f'<line x1="{ml}" y1="{yy:.1f}" x2="{ml+pw}" y2="{yy:.1f}" stroke="#eef1f5"/>')
        svg.append(f'<text x="{ml-8}" y="{yy+4:.1f}" font-size="11" fill="#9ca3af" text-anchor="end">{yv:.0f}</text>')
    svg.append(f'<line x1="{ml}" y1="{mt}" x2="{ml}" y2="{plot_bottom}" stroke="#d1d5db"/>')
    svg.append(f'<line x1="{ml}" y1="{plot_bottom}" x2="{ml+pw}" y2="{plot_bottom}" stroke="#d1d5db"/>')
    pts = [f"{xs[i]:.1f} {yp(vals[i]):.1f}" for i in range(n)]
    svg.append(f'<path d="M{" L".join(pts)}" fill="none" stroke="{color}" stroke-width="2.5"/>')
    for i in range(n):
        cx, cy = xs[i], yp(vals[i])
        y_, q, _ = series[i]
        svg.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="3.2" fill="{color}" stroke="#fff" stroke-width="1"/>')
        if q == 1:  # 仅每年 Q1 标注数值，避免 26 点拥挤
            svg.append(f'<text x="{cx:.1f}" y="{cy-10:.1f}" font-size="11" fill="#374151" text-anchor="middle">{vals[i]:.0f}</text>')
        xlab = f"{y_}Q1" if q == 1 else f"Q{q}"
        svg.append(f'<text x="{cx:.1f}" y="{plot_bottom+18:.1f}" font-size="10.5" fill="#6b7280" text-anchor="middle">{xlab}</text>')
    svg.append(f'<text x="{ml+pw}" y="{H-10}" font-size="10.5" fill="#9ca3af" text-anchor="end">单位:万吨</text>')
    svg.append('</svg>')
    with open(fname, "w", encoding="utf-8") as f:
        f.write("".join(svg))
    print(f"✅ 生成季度趋势图：{fname}（{n} 个季度点）")


# ---------- 年度趋势图（年份为 x 轴，全年合计口径，单位万吨）----------
def build_trend_y(series, fname, title, color="#c53030", unit="万吨"):
    W, H = 1000, 420
    ml, mr, mt, mb = 60, 24, 42, 56
    plot_bottom = H - mb          # 364
    ph = plot_bottom - mt         # 322
    pw = W - ml - mr              # 916
    n = len(series)
    xs = [ml + pw * i / (n - 1) for i in range(n)]
    vals = [s[1] for s in series]
    vmin, vmax = min(vals), max(vals)
    pad = (vmax - vmin) * 0.18 or 10
    vmin, vmax = vmin - pad, vmax + pad
    def yp(v):
        return mt + ph - (v - vmin) / (vmax - vmin) * ph
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" '
           f'font-family="-apple-system,Segoe UI,PingFang SC,sans-serif">']
    svg.append(f'<rect width="{W}" height="{H}" fill="#ffffff"/>')
    svg.append(f'<text x="{ml}" y="26" font-size="16" font-weight="700" fill="#1e3a5f">{title}</text>')
    for k in range(5):
        yv = vmin + (vmax - vmin) * k / 4
        yy = yp(yv)
        svg.append(f'<line x1="{ml}" y1="{yy:.1f}" x2="{ml+pw}" y2="{yy:.1f}" stroke="#eef1f5"/>')
        svg.append(f'<text x="{ml-8}" y="{yy+4:.1f}" font-size="11" fill="#9ca3af" text-anchor="end">{yv:.0f}</text>')
    svg.append(f'<line x1="{ml}" y1="{mt}" x2="{ml}" y2="{plot_bottom}" stroke="#d1d5db"/>')
    svg.append(f'<line x1="{ml}" y1="{plot_bottom}" x2="{ml+pw}" y2="{plot_bottom}" stroke="#d1d5db"/>')
    pts = [f"{xs[i]:.1f} {yp(vals[i]):.1f}" for i in range(n)]
    svg.append(f'<path d="M{" L".join(pts)}" fill="none" stroke="{color}" stroke-width="2.5"/>')
    for i in range(n):
        cx, cy = xs[i], yp(vals[i])
        y_ = series[i][0]
        svg.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="3.6" fill="{color}" stroke="#fff" stroke-width="1"/>')
        svg.append(f'<text x="{cx:.1f}" y="{cy-12:.1f}" font-size="12" fill="#374151" text-anchor="middle">{vals[i]:.0f}</text>')
        svg.append(f'<text x="{cx:.1f}" y="{plot_bottom+18:.1f}" font-size="11" fill="#6b7280" text-anchor="middle">{y_}</text>')
    svg.append(f'<text x="{ml+pw}" y="{H-10}" font-size="10.5" fill="#9ca3af" text-anchor="end">单位:{unit}</text>')
    svg.append('</svg>')
    with open(fname, "w", encoding="utf-8") as f:
        f.write("".join(svg))
    print(f"✅ 生成年度趋势图：{fname}（{n} 个年度点）")


# ---------- 卡片内 sparkline（季度）----------
def build_spark(vals, fname, color="#c53030"):
    W, H = 62, 28
    n = len(vals)
    xs = [2 + (W - 4) * i / (n - 1) for i in range(n)]
    vmin, vmax = min(vals), max(vals)
    pad = (vmax - vmin) * 0.15 or 1
    def yp(v):
        return 4 + (H - 8) - (v - (vmin - pad)) / ((vmax + pad) - (vmin - pad)) * (H - 8)
    pts = [f"{xs[i]:.1f} {yp(vals[i]):.1f}" for i in range(n)]
    svg = [f'<svg class="spark" viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg">']
    svg.append(f'<path d="M{" L".join(pts)}" fill="none" stroke="{color}" stroke-width="2"/>')
    svg.append('</svg>')
    with open(fname, "w", encoding="utf-8") as f:
        f.write("".join(svg))
    print(f"✅ 生成季度 sparkline：{fname}（{n} 点）")


# ---------- 自产/贸易煤结构（自产占比 = 自产商品煤产量 / 商品煤总销量，单位%）----------
def build_trend_ratio(series, fname, title, color="#15616d", mode="annual"):
    W, H = 1000, 420
    ml, mr, mt, mb = 60, 24, 42, 56
    plot_bottom = H - mb
    ph = plot_bottom - mt
    pw = W - ml - mr
    n = len(series)
    xs = [ml + pw * i / (n - 1) for i in range(n)]
    vals = [s[-1] for s in series]
    vmin, vmax = min(vals), max(vals)
    pad = (vmax - vmin) * 0.22 or 3
    vmin, vmax = max(0, vmin - pad), vmax + pad
    def yp(v):
        return mt + ph - (v - vmin) / (vmax - vmin) * ph
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" '
           f'font-family="-apple-system,Segoe UI,PingFang SC,sans-serif">']
    svg.append(f'<rect width="{W}" height="{H}" fill="#ffffff"/>')
    svg.append(f'<text x="{ml}" y="26" font-size="16" font-weight="700" fill="#1e3a5f">{title}</text>')
    for k in range(5):
        yv = vmin + (vmax - vmin) * k / 4
        yy = yp(yv)
        svg.append(f'<line x1="{ml}" y1="{yy:.1f}" x2="{ml+pw}" y2="{yy:.1f}" stroke="#eef1f5"/>')
        svg.append(f'<text x="{ml-8}" y="{yy+4:.1f}" font-size="11" fill="#9ca3af" text-anchor="end">{yv:.1f}%</text>')
    svg.append(f'<line x1="{ml}" y1="{mt}" x2="{ml}" y2="{plot_bottom}" stroke="#d1d5db"/>')
    svg.append(f'<line x1="{ml}" y1="{plot_bottom}" x2="{ml+pw}" y2="{plot_bottom}" stroke="#d1d5db"/>')
    pts = [f"{xs[i]:.1f} {yp(vals[i]):.1f}" for i in range(n)]
    svg.append(f'<path d="M{" L".join(pts)}" fill="none" stroke="{color}" stroke-width="2.5"/>')
    for i in range(n):
        cx, cy = xs[i], yp(vals[i])
        if mode == "annual":
            y_ = series[i][0]
            svg.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="3.6" fill="{color}" stroke="#fff" stroke-width="1"/>')
            svg.append(f'<text x="{cx:.1f}" y="{cy-12:.1f}" font-size="12" fill="#374151" text-anchor="middle">{vals[i]:.1f}%</text>')
            svg.append(f'<text x="{cx:.1f}" y="{plot_bottom+18:.1f}" font-size="11" fill="#6b7280" text-anchor="middle">{y_}</text>')
        else:
            y_, q, _ = series[i]
            svg.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="3.2" fill="{color}" stroke="#fff" stroke-width="1"/>')
            if q == 1:  # 仅每年 Q1 标注数值，避免 26 点拥挤
                svg.append(f'<text x="{cx:.1f}" y="{cy-10:.1f}" font-size="11" fill="#374151" text-anchor="middle">{vals[i]:.1f}%</text>')
            xlab = f"{y_}Q1" if q == 1 else f"Q{q}"
            svg.append(f'<text x="{cx:.1f}" y="{plot_bottom+18:.1f}" font-size="10.5" fill="#6b7280" text-anchor="middle">{xlab}</text>')
    svg.append(f'<text x="{ml+pw}" y="{H-10}" font-size="10.5" fill="#9ca3af" text-anchor="end">单位:%</text>')
    svg.append('</svg>')
    with open(fname, "w", encoding="utf-8") as f:
        f.write("".join(svg))
    print(f"✅ 生成{'年度' if mode=='annual' else '季度'}结构趋势图：{fname}（{n} 点）")


if __name__ == "__main__":
    wb = openpyxl.load_workbook(FID, data_only=True)
    sp_q = extract_quarterly(wb["自产商品煤产量"])
    ss_q = extract_quarterly(wb["商品煤总销量"])
    sp_y = extract_annual(wb["自产商品煤产量"])
    ss_y = extract_annual(wb["商品煤总销量"])
    build_trend_q(sp_q, "svg/trend_selfprod_q.svg", "中煤能源 自产商品煤产量（2020-2026·季度·万吨）", "#c53030")
    build_trend_q(ss_q, "svg/trend_totalsales_q.svg", "中煤能源 商品煤总销量（2020-2026·季度·万吨）", "#c53030")
    build_trend_y(sp_y, "svg/trend_selfprod_y.svg", "中煤能源 自产商品煤产量（2020-2025·年度·万吨）", "#c53030")
    build_trend_y(ss_y, "svg/trend_totalsales_y.svg", "中煤能源 商品煤总销量（2020-2025·年度·万吨）", "#c53030")
    build_spark([v for _, _, v in sp_q], "svg/spark_selfprod.svg", "#c53030")
    build_spark([v for _, _, v in ss_q], "svg/spark_totalsales.svg", "#c53030")
    # 自产/贸易煤结构：自产占比 = 自产商品煤产量 / 商品煤总销量
    ratio_q = [(y, q, a / b * 100) for (y, q, a), (_, _, b) in zip(sp_q, ss_q)]
    ratio_y = [(y, a / b * 100) for (y, a), (_, b) in zip(sp_y, ss_y)]
    build_trend_ratio(ratio_q, "svg/trend_structure_q.svg", "中煤能源 自产/贸易煤结构·自产占比（2020-2026·季度·%）", "#15616d", mode="q")
    build_trend_ratio(ratio_y, "svg/trend_structure_y.svg", "中煤能源 自产/贸易煤结构·自产占比（2020-2025·年度·%）", "#15616d", mode="annual")
    build_spark([v for *_, v in ratio_q], "svg/spark_structure.svg", "#15616d")
    # 在产矿井核定产能（年度，券商口径重建/待校年报）
    cap_y = extract_annual_named(wb["在产矿井核定产能(年度)"], "在产矿井核定产能")
    build_trend_y(cap_y, "svg/trend_capacity_y.svg", "中煤能源 在产矿井核定产能（2020-2025·年度·万吨）", "#d98a14")
    build_spark([v for _, v in cap_y], "svg/spark_capacity.svg", "#d98a14")
    # 剩余可采储量（年度，据年报/券商汇编；2025=129.05亿吨，可开采年限约80年）
    res_y = extract_annual_named(wb["剩余可采储量(年度)"], "剩余可采储量")
    build_trend_y(res_y, "svg/trend_reserves_y.svg", "中煤能源 剩余可采储量（2020-2025·年度·亿吨）", "#2e9e6b", unit="亿吨")
    build_spark([v for _, v in res_y], "svg/spark_reserves.svg", "#2e9e6b")

    # ============================================================
    # 价格与成本（来自 中煤能源_价格和成本.xlsx，独立 wb_pr）
    # ============================================================
    wb_pr = openpyxl.load_workbook(PR_FID, data_only=True)
    # 自产煤平均售价（年度，元/吨；红色 #c8102e 与卡片色一致）
    ap_y = extract_annual_named(wb_pr["自产煤平均售价(年度)"], "自产煤平均售价")
    build_trend_y(ap_y, "svg/trend_avgprice_y.svg", "中煤能源 自产煤平均售价（2020-2025·年度·元/吨）", "#c8102e", unit="元/吨")
    build_spark([v for _, v in ap_y], "svg/spark_avgprice.svg", "#c8102e")
    # 自产煤吨完全成本（年度，绿色 #2e9e6b）
    uc_y = extract_annual_named(wb_pr["自产煤吨完全成本(年度)"], "自产煤吨完全成本")
    build_trend_y(uc_y, "svg/trend_unitcost_y.svg", "中煤能源 自产煤吨完全成本（2020-2025·年度·元/吨）", "#2e9e6b", unit="元/吨")
    build_spark([v for _, v in uc_y], "svg/spark_unitcost.svg", "#2e9e6b")
    # 自产煤吨毛利（年度，红色 #c8102e）
    gm_y = extract_annual_named(wb_pr["自产煤吨毛利(年度)"], "自产煤吨毛利")
    build_trend_y(gm_y, "svg/trend_unitgm_y.svg", "中煤能源 自产煤吨毛利（2020-2025·年度·元/吨）", "#c8102e", unit="元/吨")
    build_spark([v for _, v in gm_y], "svg/spark_unitgm.svg", "#c8102e")
    print("✅ 完成（未写回 xlsx）")
