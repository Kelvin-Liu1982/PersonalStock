# -*- coding: utf-8 -*-
"""
卡片3「业务结构与板块协同」SVG 趋势图生成（只读数据源版）
-----------------------------------------------------------------
- 数据来源：`数据源/中煤能源_业务结构.xlsx`（6 个 sheet：产量/装备/电力 各含量与毛利率）
- 生成（每个 sheet 一张年度趋势图 + 一张 sparkline，仅年度口径）：
  - trend_biz_chem_y.svg      煤化工业务（煤化工产品总产量）
  - trend_biz_chem_gm_y.svg   煤化工业务毛利率（分部毛利率 %）
  - trend_biz_equip_y.svg     煤矿装备业务（工业产值）
  - trend_biz_equip_gm_y.svg  煤矿装备业务毛利率（分部毛利率 %）
  - trend_biz_elec_y.svg      电力业务（年度发电量）
  - trend_biz_elec_gm_y.svg   电力业务毛利率（分部毛利率 %）
  - spark_biz_chem.svg / spark_biz_chem_gm.svg / spark_biz_equip.svg / spark_biz_equip_gm.svg / spark_biz_elec.svg / spark_biz_elec_gm.svg
- 视觉风格与 build_finance_svg.py 的 build_svg_base / build_sparkline 完全一致。
- 本脚本**只读 xlsx，绝不写回 Excel**。
"""
import openpyxl

EXCEL = "数据源/中煤能源_业务结构.xlsx"

# 6 个指标定义：key, 标题, sheet名, 列名, 单位, 颜色
METRICS = [
    ("biz_chem",      "煤化工产品总产量", "煤化工业务年度",       "煤化工产品总产量",     "万吨", "#1e3a5f"),
    ("biz_chem_gm",   "煤化工毛利率",     "煤化工业务年度毛利率", "分部毛利率",           "%",   "#c53030"),
    ("biz_equip",     "煤矿装备工业产值", "煤矿装备业务年度",     "煤矿装备工业产值",     "亿元", "#d97706"),
    ("biz_equip_gm",  "煤矿装备毛利率",   "煤矿装备业务年度毛利率","分部毛利率",          "%",   "#b45309"),
    ("biz_elec",      "电力发电量",       "电力业务年度",         "年度发电量",           "亿千瓦时", "#047857"),
    ("biz_elec_gm",   "电力毛利率",       "电力业务年度毛利率",   "分部毛利率",           "%",   "#2563eb"),
]

def parse_val(v):
    """把单元格值解析为数值或 None（容忍 '-' / '—' / '未单独披露' / '约 X' / 百分比字符串）"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("-", "—", "未单独披露", "", "None"):
        return None
    if s.startswith("约"):
        s = s[1:]
    if s.endswith("%"):
        try:
            return float(s[:-1])
        except ValueError:
            return None
    try:
        return float(s)
    except ValueError:
        return None

# ---------- 读取 sheet -> (labels, vals) ----------
def load_col(sheet_name, col_name):
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb[sheet_name]
    hdr = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    ci = hdr.get(col_name)
    labels, vals = [], []
    import re as _re
    for r in range(2, ws.max_row + 1):
        yv = ws.cell(row=r, column=1).value
        if yv is None:
            continue
        # 仅接受纯年份标签（'2020' 或 '2020年'），过滤掉说明性文字行
        # 不允许包含中文标点 / 字母 / 多于 1 个连续数字段
        s = str(yv).strip()
        if not _re.fullmatch(r"\d{4}年?", s):
            continue
        ystr = s.replace("年", "")
        labels.append(ystr)
        cv = ws.cell(row=r, column=ci).value if ci else None
        vals.append(parse_val(cv))
    return labels, vals

# ---------- 绘图函数（与 build_finance_svg.py 视觉一致） ----------
def build_svg_base(series_list, title, yunit, fname, W=1000, H=420, xlabels=None):
    ml, mr, mt, mb = 58, 24, 42, 44
    pw, ph = W - ml - mr, H - mt - mb
    n = len(series_list[0][1]) if series_list else 1
    xs = [ml + (pw * i / (n - 1)) for i in range(n)]
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" font-family="-apple-system,Segoe UI,PingFang SC,sans-serif">'.format(W=W, H=H)]
    svg.append(f'<rect width="{W}" height="{H}" fill="#ffffff"/>')
    svg.append(f'<text x="{ml}" y="26" font-size="16" font-weight="700" fill="#1e3a5f">{title}</text>')
    allv = [v for _, vals, _ in series_list for v in vals if v is not None]
    ymin, ymax = min(allv), max(allv)
    pad = (ymax - ymin) * 0.15 or 1
    ymin, ymax = ymin - pad, ymax + pad
    def ypos(v):
        return mt + ph - (v - ymin) / (ymax - ymin) * ph
    for k in range(5):
        yv = ymin + (ymax - ymin) * k / 4
        yy = ypos(yv)
        svg.append(f'<line x1="{ml}" y1="{yy:.1f}" x2="{ml+pw}" y2="{yy:.1f}" stroke="#eef1f5"/>')
        svg.append(f'<text x="{ml-8}" y="{yy+4:.1f}" font-size="11" fill="#9ca3af" text-anchor="end">{yv:.1f}</text>')
    svg.append(f'<line x1="{ml}" y1="{mt}" x2="{ml}" y2="{mt+ph}" stroke="#d1d5db"/>')
    svg.append(f'<line x1="{ml}" y1="{mt+ph}" x2="{ml+pw}" y2="{mt+ph}" stroke="#d1d5db"/>')
    for i in range(n):
        _lab = (xlabels[i] if xlabels else str(i))
        svg.append(f'<text x="{xs[i]:.1f}" y="{mt+ph+18:.1f}" font-size="10.5" fill="#6b7280" text-anchor="middle">{_lab}</text>')
    for name, vals, color in series_list:
        pts = [(xs[i], ypos(vals[i])) for i in range(n) if vals[i] is not None]
        if len(pts) > 1:
            d = "M" + " L".join(f"{x:.1f} {y:.1f}" for x, y in pts)
            svg.append(f'<path d="{d}" fill="none" stroke="{color}" stroke-width="2.5"/>')
        for x, y in pts:
            svg.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3" fill="{color}" stroke="#fff" stroke-width="1"/>')
    lx = ml + 4
    for name, _, color in series_list:
        svg.append(f'<rect x="{lx}" y="{H-22}" width="12" height="12" rx="2" fill="{color}"/>')
        svg.append(f'<text x="{lx+17}" y="{H-12}" font-size="11.5" fill="#374151">{name}</text>')
        lx += 18 + len(name) * 12
    svg.append(f'<text x="{ml+pw}" y="{H-12}" font-size="10.5" fill="#9ca3af" text-anchor="end">单位:{yunit}</text>')
    svg.append('</svg>')
    s = "".join(svg)
    with open(fname, "w", encoding="utf-8") as f:
        f.write(s)
    print(f"OK {fname}")
    return s

def build_sparkline(vals, color, fname):
    W, H = 62, 28
    n = len(vals)
    clean = [v for v in vals if v is not None]
    if len(clean) < 2:
        d = f"M2 {H/2:.1f} L{W-2} {H/2:.1f}"
    else:
        ymin, ymax = min(clean), max(clean)
        pad = (ymax - ymin) * 0.1 or 1
        ymin, ymax = ymin - pad, ymax + pad
        pts = []
        for i, v in enumerate(vals):
            if v is None:
                continue
            x = 2 + (W - 4) * i / (n - 1)
            y = H - 2 - (v - ymin) / (ymax - ymin) * (H - 4)
            pts.append(f"{x:.1f} {y:.1f}")
        d = "M" + " L".join(pts)
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" class="spark">'
           f'<path d="{d}" fill="none" stroke="{color}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"/></svg>')
    with open(fname, "w", encoding="utf-8") as f:
        f.write(svg)
    print(f"OK {fname}")
    return svg

for key, title, sheet, col, unit, color in METRICS:
    labels, vals = load_col(sheet, col)
    build_svg_base([(title, vals, color)], f"{title}趋势（年度）", unit, f"trend_{key}_y.svg", W=1000, H=420, xlabels=labels)
    build_sparkline(vals, color, f"spark_{key}.svg")

print("\nDone: 6 张年度趋势图 + 6 张 sparkline 已生成（仅年度口径，均来自 中煤能源_业务结构.xlsx）")
