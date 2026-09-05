# -*- coding: utf-8 -*-
"""中国电信 · 财务健康与盈利质量：XLSX(年报/季报) → 静态 SVG。

读取「中国电信_财务健康与盈利质量.xlsx」：
  - 年报数据（年份为行、指标为列）→ 每张指标 1 个年度 SVG
  - 季报数据（季度为行、指标为列）→ 每张指标 1 个季度 SVG
年度-only 指标（季报无对应列）→ 生成占位 SVG。

输出：svg/s1/s1_<key>_annual.svg 与 s1_<key>_q.svg
只读 XLSX，绝不写回。

基于 finance-xlsx-svg-charts skill 的紧凑边距折线图（现由 common.svg_charts.make_line_svg 提供，与 china-mobile 看板一致）。
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from common.svg_charts import make_line_svg, make_placeholder

import openpyxl

XLSX = os.path.join(HERE, "数据源", "中国电信_财务健康与盈利质量.xlsx")
OUT = os.path.join(HERE, "svg", "s1")
os.makedirs(OUT, exist_ok=True)

# (表头关键字, key, 显示标题, 单位, 颜色)
ANNUAL = [
    ("营业收入", "revenue", "营业收入", "亿元", "#c8102e"),
    ("归母净利润", "net_profit", "归母净利润", "亿元", "#15616d"),
    ("EBITDA", "ebitda", "EBITDA", "亿元", "#2e9e6b"),
    ("EBITDA率", "ebitda_rate", "EBITDA率", "%", "#2e9e6b"),
    ("ROE", "roe", "净资产收益率 ROE", "%", "#d98a14"),
    ("净利率", "net_margin", "净利率", "%", "#7a8694"),
    ("资产负债率", "debt_ratio", "资产负债率", "%", "#2c7fb8"),
    ("自由现金流", "fcf", "自由现金流", "亿元", "#6a4c93"),
    ("应收账款", "ar", "应收账款", "亿元", "#b5651d"),
]
QUARTERLY = [
    ("营业收入", "revenue", "营业收入(单季)", "亿元", "#c8102e"),
    ("归母净利润", "net_profit", "归母净利润(单季)", "亿元", "#15616d"),
    ("EBITDA", "ebitda", "EBITDA(单季)", "亿元", "#2e9e6b"),
    ("EBITDA率", "ebitda_rate", "EBITDA率(单季)", "%", "#2e9e6b"),
    ("ROE", "roe", "单季 ROE", "%", "#d98a14"),
    ("净利率", "net_margin", "净利率(单季)", "%", "#7a8694"),
    ("资产负债率", "debt_ratio", "资产负债率(季末)", "%", "#2c7fb8"),
    ("应收账款", "ar", "应收账款(季末)", "亿元", "#b5651d"),
]


def col_of(ws, header_row, keyword):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is not None and keyword in str(v):
            return c
    return None


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("", "—", "-", "未单独披露", "约", "N/A", "None"):
        return None
    try:
        return float(s.replace("pp", "").replace("+", "").replace("%", "").replace(",", ""))
    except Exception:
        return None


wb = openpyxl.load_workbook(XLSX, data_only=True)

# ---------- 年度 ----------
ws = wb["年报数据"]
HR = 3
for kw, key, title, unit, color in ANNUAL:
    c = col_of(ws, HR, kw)
    if c is None:
        print("SKIP annual", key, "(列未找到)")
        continue
    labels, vals = [], []
    for r in range(HR + 1, ws.max_row + 1):
        yv = ws.cell(row=r, column=1).value
        if yv is None:
            continue
        labels.append(int(yv) if str(yv).isdigit() else yv)
        vals.append(num(ws.cell(row=r, column=c).value))
    svg = make_line_svg(title, unit, labels, vals, color, W=480, H=260, quarterly=False)
    fn = os.path.join(OUT, f"s1_{key}_annual.svg")
    with open(fn, "w", encoding="utf-8") as f:
        f.write(svg)
    print("wrote", fn, "pts=%d" % len(vals))

# ---------- 季度 ----------
ws = wb["季报数据"]
HR = 3
for kw, key, title, unit, color in QUARTERLY:
    c = col_of(ws, HR, kw)
    if c is None:
        print("SKIP quarterly", key, "(列未找到)")
        continue
    labels, vals = [], []
    for r in range(HR + 1, ws.max_row + 1):
        qv = ws.cell(row=r, column=1).value
        if qv is None:
            continue
        qv = str(qv)
        if "Q" not in qv:
            continue
        labels.append(qv)
        vals.append(num(ws.cell(row=r, column=c).value))
    svg = make_line_svg(title, unit, labels, vals, color, W=760, H=260, quarterly=True)
    fn = os.path.join(OUT, f"s1_{key}_q.svg")
    with open(fn, "w", encoding="utf-8") as f:
        f.write(svg)
    print("wrote", fn, "pts=%d" % len(vals))

# ---------- 年度-only 指标占位（季报无对应列）----------
q_keys = {k for _, k, _, _, _ in QUARTERLY}
for kw, key, title, unit, color in ANNUAL:
    if key not in q_keys:
        svg = make_placeholder(title, 760, 260)
        fn = os.path.join(OUT, f"s1_{key}_q.svg")
        with open(fn, "w", encoding="utf-8") as f:
            f.write(svg)
        print("wrote placeholder", fn)
