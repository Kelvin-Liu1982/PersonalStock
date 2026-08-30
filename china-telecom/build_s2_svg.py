# -*- coding: utf-8 -*-
"""中国电信 · 核心业务运营与结构转型：年报数据(指标为行) → 年度 SVG。

源 xlsx 布局：年报数据 sheet，row4 = 表头(指标,2020..2025)，row5+ = 各指标年度值。
仅含年报数据，故只生成年度 SVG 到 svg/s2/。生成器只读，不写回 xlsx。
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_line_svg import make_line_svg

SRC = "数据源/中国电信_核心业务运营与结构转型.xlsx"
OUT = "svg/s2"
os.makedirs(OUT, exist_ok=True)

# 指标名 -> (svg key, 折线颜色)
SPEC = {
    "服务收入(亿元)": ("service_rev", "#c8102e"),
    "产业数字化收入(亿元)": ("ind_digital", "#15616d"),
    "天翼云收入(亿元)": ("tianyi_cloud", "#c8102e"),
    "IDC收入(亿元)": ("idc", "#2e9e6b"),
    "智能(AI)收入(亿元)": ("ai", "#d98a14"),
    "安全收入(亿元)": ("security", "#2e9e6b"),
}

PLACEHOLDER = {"", "—", "-", "'-'", "未单独披露", "约", "待补"}


def to_val(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s in PLACEHOLDER:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return float(v)


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["年报数据"]
    labels = [str(ws.cell(row=4, column=c).value) for c in range(2, 8)]  # 2020..2025
    print("labels:", labels)

    for r in range(5, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or name not in SPEC:
            if name:
                print("  skip (no spec):", name)
            continue
        key, color = SPEC[name]
        raw = [ws.cell(row=r, column=c).value for c in range(2, 8)]
        vals = [to_val(v) for v in raw]
        title = name.replace("(亿元)", "")
        svg = make_line_svg(title, "亿元", labels, vals, color)
        fn = os.path.join(OUT, f"s2_{key}_annual.svg")
        with open(fn, "w", encoding="utf-8") as f:
            f.write(svg)
        npts = sum(1 for v in vals if v is not None)
        print(f"  wrote {fn}  points={npts}  vals={vals}")


if __name__ == "__main__":
    main()
